#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把脱敏 Markdown 转成可上传的 .docx。
处理：# 一级标题、## 二级标题、表格、普通段落。
用法: python md2docx.py <输入.md> <输出.docx>
"""
import re
import sys
from docx import Document
from docx.shared import Pt, RGBColor


def add_paragraph_or_table(doc, lines):
    """取一段连续的（普通段落行 或 表格块）写入 docx。"""
    # 检测是否整块为 markdown 表格（含 | 分隔行）
    if all(re.search(r"\|", ln) for ln in lines if ln.strip()):
        rows = []
        for ln in lines:
            ln = ln.strip()
            if not ln or re.match(r"^\|[\s\-:|]+\|?$", ln):
                continue  # 跳过空行与分隔行
            cells = [c.strip() for c in ln.strip("|").split("|")]
            rows.append(cells)
        if rows:
            tbl = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            tbl.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci in range(len(row)):
                    tbl.cell(ri, ci).text = row[ci]
            return
    # 普通段落
    for ln in lines:
        if ln.strip():
            doc.add_paragraph(ln.strip())


def md2docx(md_path, out_path):
    doc = Document()
    buf = []
    with open(md_path, encoding="utf-8") as f:
        for raw in f:
            line = raw.rstrip("\n")
            s = line.strip()
            # 处理标题
            m = re.match(r"^(#{1,6})\s+(.*)$", s)
            if m:
                if buf:
                    add_paragraph_or_table(doc, buf)
                    buf = []
                level = len(m.group(1))
                text = m.group(2)
                h = doc.add_heading(text, level=min(level, 4))
                for run in h.runs:
                    run.font.color.rgb = RGBColor(0, 0, 0)
                continue
            # 忽略图片引用行
            if re.match(r"!\[.*\]\(.*\)", s):
                continue
            buf.append(line)
    if buf:
        add_paragraph_or_table(doc, buf)
    doc.save(out_path)
    print("已生成:", out_path)


if __name__ == "__main__":
    md2docx(sys.argv[1], sys.argv[2])
